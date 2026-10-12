import os
import shutil
import docx
from docx.shared import Inches
from pptx import Presentation
from pptx.util import Inches as PptxInches
from PIL import Image

DESKTOP = r"C:\Users\satis\OneDrive\Desktop"
PROJECT_DIR = r"c:\Users\satis\OneDrive\Desktop\movie-recommender"

STD_IMG = os.path.join(PROJECT_DIR, "reports", "figures", "dfd_level_1.png")
CHAT_IMG = os.path.join(PROJECT_DIR, "reports", "figures_chatbot", "dfd_level_1.png")

def update_docx(filename, new_img_path):
    filepath = os.path.join(DESKTOP, filename)
    print(f"\n[Updating DOCX] {filename} with {new_img_path}")
    
    # Backup
    backup_path = filepath + ".bak"
    shutil.copy2(filepath, backup_path)
    print(f"  Backup created: {backup_path}")
    
    with open(new_img_path, "rb") as f:
        new_bytes = f.read()
        
    with Image.open(new_img_path) as im:
        orig_w, orig_h = im.size
        aspect = orig_w / orig_h
        
    doc = docx.Document(filepath)
    
    # Update DFD1 part
    part = doc.part.related_parts["rId9"]
    old_size = len(part.blob)
    part._blob = new_bytes
    print(f"  Updated image part {part.partname}: old={old_size} bytes, new={len(part.blob)} bytes")
    
    # Update shape dimensions to maintain pristine aspect ratio at width 6.2 inches
    target_w = Inches(6.2)
    target_h = int(target_w / aspect)
    updated_shapes = 0
    for s in doc.inline_shapes:
        if "rId9" in s._inline.xml:
            s.width = target_w
            s.height = target_h
            updated_shapes += 1
            print(f"  Adjusted inline shape: width={s.width}, height={s.height}")
            
    doc.save(filepath)
    print(f"  Saved {filename} successfully ({updated_shapes} shapes adjusted).")

def update_pptx(filename, slide_index, new_img_path):
    filepath = os.path.join(DESKTOP, filename)
    print(f"\n[Updating PPTX] {filename} (Slide {slide_index + 1}) with {new_img_path}")
    
    # Backup
    backup_path = filepath + ".bak"
    shutil.copy2(filepath, backup_path)
    print(f"  Backup created: {backup_path}")
    
    with open(new_img_path, "rb") as f:
        new_bytes = f.read()
        
    with Image.open(new_img_path) as im:
        orig_w, orig_h = im.size
        aspect = orig_w / orig_h
        
    prs = Presentation(filepath)
    slide = prs.slides[slide_index]
    
    pics = [s for s in slide.shapes if s.shape_type == 13]
    if not pics:
        raise RuntimeError(f"No picture shape found on slide {slide_index + 1}")
        
    pic = pics[0]
    rId = pic._element.xpath(".//a:blip/@r:embed")[0]
    target_part = slide.part.related_part(rId)
    old_size = len(target_part.blob)
    target_part._blob = new_bytes
    print(f"  Updated image part {target_part.partname}: old={old_size} bytes, new={len(target_part.blob)} bytes")
    
    # Calculate fitting dimensions matching slide template
    max_w = 11.0
    max_h = 5.0
    if (max_w / max_h) > aspect:
        final_h = max_h
        final_w = final_h * aspect
    else:
        final_w = max_w
        final_h = final_w / aspect
        
    left = (13.333 - final_w) / 2.0
    top = 1.65 + (5.1 - final_h) / 2.0
    
    pic.width = PptxInches(final_w)
    pic.height = PptxInches(final_h)
    pic.left = PptxInches(left)
    pic.top = PptxInches(top)
    print(f"  Adjusted picture bounds: width={final_w:.3f}\", height={final_h:.3f}\", left={left:.3f}\", top={top:.3f}\"")
    
    prs.save(filepath)
    print(f"  Saved {filename} successfully.")

if __name__ == "__main__":
    # 1. Standard Report
    update_docx("Movie_Recommendation_System_Project_Report.docx", STD_IMG)
    # 2. Chatbot Report
    update_docx("Movie_Recommender_Chatbot_Project_Report.docx", CHAT_IMG)
    # 3. Standard Presentation (Slide 6 = index 5)
    update_pptx("Movie_Recommendation_System_Presentation.pptx", 5, STD_IMG)
    # 4. Chatbot Presentation (Slide 8 = index 7)
    update_pptx("Movie_Recommender_Chatbot_Presentation.pptx", 7, CHAT_IMG)
    print("\nAll 4 files updated successfully!")
